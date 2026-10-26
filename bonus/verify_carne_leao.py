"""Checks the workbook formulas against the Python model, with an independent formula engine.

    python verify_carne_leao.py

1. Hand-computed tax cases for the 2026 rules.
2. Builds a filled test copy (every branch: exempt, partial reduction, full table, legal
   deductions, < R$ 10 carry-over, wrong year, missing data, missing rate), recalculates it
   with the `formulas` engine and compares every formula cell with the model's value.
3. Does the same for the blank template that is delivered.
"""

import datetime as dt
import sys
import tempfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import formulas
import formulas.functions.math as formulas_math
import numpy as np
from formulas.functions import XlError, wrap_ufunc
import openpyxl

import build_carne_leao as b

D = dt.date

# The `formulas` engine implements WORKDAY with numpy's busday_offset in its default strict
# mode, which returns #VALUE! when the start date is a weekend or holiday. Excel (and
# LibreOffice, Google Sheets) accept any start date: going back from a non-business day
# lands on the previous business day. Match Excel so the due-date cells can be checked.
_busday_offset = np.busday_offset


def _excel_busday_offset(dates, offsets, roll=None, **kw):
    if roll is None:
        roll = "forward" if np.all(np.asarray(offsets) <= 0) else "backward"
    return _busday_offset(dates, offsets, roll=roll, **kw)


np.busday_offset = _excel_busday_offset


# Same for ROUND: the engine rounds the raw float (1702.6149999999998 -> 1702.61), while
# Excel first reduces to 15 significant digits (1702.615 -> 1702.62).
def _excel_round(x, d):
    step = Decimal(1).scaleb(-int(d))
    return float(Decimal(format(float(x), ".15g")).quantize(step, rounding=ROUND_HALF_UP))


formulas_math.FUNCTIONS["ROUND"] = wrap_ufunc(_excel_round)


def check_rules():
    def month_tax(gross, legal=0.0):
        base = max(0.0, gross - max(legal, b.SIMPLIFIED))
        tax = b.table_tax(base)
        return b.round2(tax - b.reduction(gross, tax))

    cases = [  # gross (R$), expected tax, worked by hand
        (5000.00, 0.00),     # (5000 - 607,20) x 22,5% - 675,49 = 312,89, fully reduced
        (6000.00, 394.54),   # 5392,80 x 27,5% - 908,73 = 574,29; reduction 978,62 - 798,87 = 179,75
        (7350.00, 945.54),   # 6742,80 x 27,5% - 908,73 = 945,54; reduction rounds to 0,00
        (8000.00, 1124.29),  # 7392,80 x 27,5% - 908,73; no reduction
        (2000.00, 0.00),
    ]
    ok = True
    for gross, expected in cases:
        got = month_tax(gross)
        ok &= abs(got - expected) < 0.005
        print(f"  R$ {gross:>9,.2f} -> {got:>8,.2f} (expected {expected:,.2f})")
    return ok


TEST_PAYMENTS = [
    (D(2026, 1, 5), "Outlier", 480.00),
    (D(2026, 2, 10), "DataAnnotation", 900.00),
    (D(2026, 3, 3), "Outlier", 650.00),
    (D(2026, 3, 20), "Mercor", 500.00),
    (D(2026, 4, 8), "Alignerr", 1500.00),
    (D(2026, 5, 12), "Outlier", 3000.00),
    (D(2026, 6, 9), "Outlier", 989.20),
    (D(2026, 8, 14), "DataAnnotation", 993.70),
    (D(2026, 9, 2), "Outlier", 1500.00),
    (D(2026, 10, 1), "Outlier", 400.00),
    (D(2026, 10, 15), "Mercor", 350.50),
    (D(2025, 12, 20), "Outlier", 100.00),
    (D(2026, 10, 20), "Mercor", None),
    (None, "Outlier", 50.00),
    (D(2026, 11, 3), "Outlier", 2000.00),
    (D(2026, 12, 4), "Outlier", 700.00),
]
TEST_DEDUCTIONS = {9: (1000.00, 2, 0, 0), 11: (0, 0, 300.00, 0)}


def compare(path: Path) -> tuple[int, list[str]]:
    """Engine result vs cached (model) value for every formula cell."""
    xl = formulas.ExcelModel().loads(str(path)).finish()
    sol = xl.calculate()
    engine = {}
    for key, rng in sol.items():
        if "!" not in key:
            continue
        sheet, ref = key.rsplit("!", 1)
        sheet = sheet.strip("'").split("]", 1)[1]
        value = rng.value[0, 0] if hasattr(rng, "value") else rng
        engine[(sheet, ref)] = value

    wf = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)
    checked, problems = 0, []
    for ws in wf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not (isinstance(c.value, str) and c.value.startswith("=")):
                    continue
                checked += 1
                cached = wv[ws.title][c.coordinate].value
                got = engine.get((ws.title.upper(), c.coordinate))
                if isinstance(got, XlError) or "Error" in type(got).__name__:
                    problems.append(f"{ws.title}!{c.coordinate}: engine error {got}")
                    continue
                if hasattr(got, "item"):
                    got = got.item()
                if got is None or (not isinstance(got, (int, float, str)) and str(got) == ""):
                    got = ""
                if isinstance(cached, dt.datetime):
                    cached = b.excel_date(cached.date())
                if cached is None:
                    cached = ""
                same = (abs(float(got) - float(cached)) < 0.005
                        if isinstance(got, (int, float)) and isinstance(cached, (int, float))
                        else str(got) == str(cached))
                if not same:
                    problems.append(f"{ws.title}!{c.coordinate}: engine={got!r} model={cached!r}")
    return checked, problems


def main():
    ok = True
    print("1. Tax rules, hand-computed cases")
    ok &= check_rules()

    with tempfile.TemporaryDirectory() as tmp:
        test_file = Path(tmp) / "teste.xlsx"
        b.build(test_file, TEST_PAYMENTS, TEST_DEDUCTIONS)
        checked, problems = compare(test_file)
        print(f"2. Filled test copy: {checked} formulas checked, {len(problems)} mismatches")
        for p in problems[:30]:
            print("   ", p)
        ok &= not problems

        wv = openpyxl.load_workbook(test_file, data_only=True)["Resumo do mês"]
        print("   Month | Recebido | Imposto do mês | Saldo ant. | DARF | Vencimento | Situação")
        for r in range(6, 18):
            v = [wv[f"{c}{r}"].value for c in "CDIJKLM"]
            print(f"   {v[0]:<9} {v[1]:>10,.2f} {v[2]:>10,.2f} {v[3]:>8,.2f} {v[4]:>9,.2f}  "
                  f"{v[5]:%d/%m/%Y}  {v[6]}")

    checked, problems = compare(b.OUT)
    print(f"3. Delivered template: {checked} formulas checked, {len(problems)} mismatches")
    for p in problems[:30]:
        print("   ", p)
    ok &= not problems

    print("ALL CHECKS PASSED" if ok else "CHECKS FAILED")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
